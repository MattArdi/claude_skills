#!/usr/bin/env python3
"""Render a contract review into the standard review workbook.

Usage:
  build_review_workbook.py <standards.json> <review.json> <out.xlsx>
  build_review_workbook.py <standards.json> --template <out.xlsx>

review.json (all keys optional):
  summary:  contract, file, status (Draft|Executed), contract_type (Service|Funding|Vendor|Other),
            paper, standards_version, our_entity, our_role (Provider|Buyer|Other), role_basis, counterparty,
            defined_terms ({"Start2": "the Contractor", "Company": "Canada"}: the contract's names for the parties)
  facts:    [{fact, value, where}]
  findings: [{severity, type, clause, where, quote, issue, action, subjective, suggested_language}]
  coverage: {CLAUSE_ID: {verdict, where, note}}
  checks:   {"1": {result, evidence}, ...}   (keys are Execution Check numbers)
  documents_not_supplied: [{document, referenced_at, quote, part_of_contract, governs, effect, action}]
Quotes may mark the words a finding turns on with **bold**; they are rendered bold in Excel.
Suggested language is the finding's own suggested_language, else the Start2 standard language of a standard clause
(Missing, Partial or Deviant findings) with the standard's defined terms swapped for the contract's.
Counts and the to-confirm line are formulas over the Findings sheet, so edits in Excel flow through.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from suggest_language import render  # noqa: E402

SEV = ["Critical", "High", "Medium", "Low"]
TYPES = ["Missing", "Partial", "Deviant", "Client: Reject", "Client: Escalate", "Client: Conditional",
         "Client: Accept", "Unclassified clause", "Commercial", "Execution", "Not supplied", "Other"]
VERDICTS = ["Conforming", "Partial", "Deviant", "Missing", "Not supplied", "N/A"]
RESULTS = ["Pass", "Fail", "N/A", "Cannot verify"]
PAPER = ["Start2 template, unmodified", "Start2 template, modified", "Client or third-party paper"]
SIGNOFF = ("1. Director of People and Operations; 2. Regional Director of Innovation and Product Strategy; 3. CEO. "
           "Applies to every contract.")
CHECK_LABEL_OVERRIDE = {11: "Cross-references in the contract point to clauses and annexes that exist and are correctly numbered."}

HFILL = PatternFill("solid", fgColor="FF1F3864"); HFONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
SFILL = PatternFill("solid", fgColor="FFD9E1F2"); BFONT = Font(name="Arial", size=10); BOLD = Font(name="Arial", size=10, bold=True)
thin = Side(style="thin", color="FFBFBFBF"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
COLORS = {"Critical": "FFF4B6B6", "High": "FFF9D6A5", "Medium": "FFFFF2B3", "Low": "FFDCEBF7"}
MAXROW = 300
DOCS = "Documents not supplied"
PLAIN, STRONG = InlineFont(rFont="Arial", sz=10), InlineFont(rFont="Arial", sz=10, b=True)


def rich(text):
    """Render **bold** markers as bold runs; plain strings pass through."""
    if not isinstance(text, str) or "**" not in text:
        return text
    return CellRichText([TextBlock(STRONG if i % 2 else PLAIN, part) for i, part in enumerate(text.split("**")) if part])


def suggestion(f, by, terms):
    if f.get("suggested_language"):
        return f["suggested_language"]
    if f.get("type") not in ("Missing", "Partial", "Deviant"):
        return None
    text = render(by, f.get("clause"), terms)
    if not text:
        return None
    head = f"Start2 standard clause {f['clause']} ({by[f['clause']]['draft_heading']}), worded for this contract:"
    tail = "\n\nFill the bracketed items before sending." if re.search(r"\[[^\]]+\]", text) else ""
    return f"{head}\n{text}{tail}"


def cell(ws, r, c, v=None, font=BFONT, fill=None, wrap=True):
    x = ws.cell(r, c, v)
    x.font = font; x.border = BORDER
    x.alignment = Alignment(wrap_text=wrap, vertical="top")
    if fill: x.fill = fill
    return x


def header(ws, heads, widths):
    for j, (h, w) in enumerate(zip(heads, widths), 1):
        c = cell(ws, 1, j, h, HFONT, HFILL); c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def colour(ws, rng):
    for k, col in COLORS.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", bgColor=col, fgColor=col)))


def dropdown(ws, rng, items):
    dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(rng)


def readme(ws):
    rows = [
        ("S", "What this workbook is"),
        ("One contract per workbook", "Produced by the contract review skill from the contract text and the Start2 standards workbook, then checked by a person. Counts and the to-confirm line update when you edit the Findings sheet."),
        ("S", "How to use it"),
        ("1. Summary", "Read the role, the counts and the number of findings to confirm. Check the role guess first: if it is wrong, nothing else applies."),
        ("2. Findings", "Work down the list. It is sorted Critical to Low. Fix, negotiate or accept each item, and change Severity or Type if you disagree. Every Critical, High and client-clause finding must be confirmed by a person."),
        ("3. Coverage", "Use it to see which standard clauses are present, partial, deviant or missing. It is a view; issues are listed in Findings."),
        ("4. Checks", "The 12 execution checks. A failed check also appears in Findings."),
        ("5. Documents not supplied", "Every outside document the contract names (general conditions, bids, policies, codes, links) that was not supplied with it. Terms in these documents were not reviewed. Obtain them and re-run the review."),
        ("S", "Severity"),
        ("Critical", "Do not sign as it stands. Reject threshold met, hard stop, wrong entity, no liability cap, blank essential fields."),
        ("High", "Fix, or have the full sign-off chain knowingly accept it, before signing. Material cash, liability, IP, data, forum or scope exposure."),
        ("Medium", "Should be fixed. Negotiable. Raises dispute risk but not direct cash loss."),
        ("Low", "Housekeeping, optional clauses, accepted client clauses."),
        ("Subjective wording", "A finding whose quoted text lets the counterparty alone judge Start2's work, payment or people (for example 'acceptable to Canada', 'to the satisfaction of', 'in its sole discretion') is raised one severity level and marked Yes in the Subjective column."),
        ("S", "Sheet: Summary (Item / Value / Where)"),
        ("Contract", "Name or title of the contract being reviewed."),
        ("Source file", "File name of the contract that was reviewed."),
        ("Status", "Draft (unsigned) or Executed (signed). Executed contracts get remediation wording instead of sign/no-sign."),
        ("Contract type", "Service, Funding, Vendor or Other. Funding agreements (grants, claims, audits) read the fee and cost clauses differently."),
        ("Whose paper", "Start2 template unmodified, Start2 template modified, or client or third-party paper. Sets the review lead time: 3 business days for the unmodified template, otherwise 5."),
        ("Standards version", "Version of the standards workbook the contract was checked against."),
        ("Our entity", "The Start2 entity in this contract. Guessed from the entity list and the text when not obvious."),
        ("Our role", "Provider, Buyer or Other. The standards cover provider-side contracts only. If the role is not Provider, standards coverage is not applied and the To confirm line says so."),
        ("Role basis", "One sentence on why that role was chosen. Always filled when the role was guessed."),
        ("Counterparty", "The other party."),
        ("Critical / High / Medium / Low", "Formulas: number of Findings at each severity."),
        ("To confirm", "Formula: the number of Critical and High findings. These are logged for a person to confirm; nothing else is recommended."),
        ("Documents not supplied", "Formula: the number of rows on the Documents not supplied sheet."),
        ("Review lead time", "Formula from Whose paper."),
        ("Sign-off chain", "Fixed. Every contract goes to all three approvers in order."),
        ("Key facts (Fact / Value / Where)", "Facts used in the review: parties, dates, amounts, law, forum, language, annexes, with arithmetic checks. Where gives the contract reference and page."),
        ("S", "Sheet: Findings (one row per issue)"),
        ("#", "Running number."),
        ("Severity", "Critical, High, Medium or Low (drop-down)."),
        ("Type", "Missing, Partial or Deviant (a standard clause); Client: Reject, Escalate, Conditional or Accept (a client-added clause matched to the playbook); Unclassified clause (not in the standard or playbook, generic screen); Commercial (fee, dates, arithmetic); Execution (a failed check); Not supplied (a document that forms part of the contract was not supplied, so its terms were not reviewed); Other."),
        ("Clause", "Standard clause ID (for example LOL), the playbook row name, or Check n."),
        ("Where", "Contract clause number and page, for example 13.1 (p. 10). 'Not in contract' for an absent clause."),
        ("Quote", "Verbatim text from the contract: the full sentence or sub-clause, with the words the finding turns on in bold. '...' marks text left out between excerpts. Scanned contracts are transcribed from page images, so check the page before relying on a quote."),
        ("Issue", "What the clause says in plain words, what it means for Start2 in practice, how it compares with the Start2 standard, and why it has this severity."),
        ("Action", "What to do next and who does it: what to check internally, what to ask the counterparty for, and a fallback if they decline."),
        ("Suggested language", "Wording to propose to the counterparty. For a standard clause, the Start2 standard language with the parties named as in this contract; otherwise wording written for this finding. Fill bracketed items and adjust before sending."),
        ("Subjective", "Yes when the quoted text lets the counterparty alone judge Start2's work, payment or people. The severity already includes the one-level raise."),
        ("S", "Sheet: Coverage (one row per standard clause)"),
        ("ID", "Stable clause key from the standards workbook."),
        ("Clause", "Clause name."),
        ("Base severity", "Severity if the clause is missing or deviant. Editable for this contract (for example raise anti-bribery to Medium for a government counterparty). n/a for internal rules."),
        ("Verdict", "Conforming, Partial, Deviant, Missing, Not supplied (the contract says the topic is governed by a document that was not supplied), or N/A (drop-down)."),
        ("Where", "Contract clause number and page where the clause was found."),
        ("Severity", "Formula. Missing, Deviant or Not supplied = Base severity; Partial = one level lower (floor Low); Conforming or N/A = '-'."),
        ("Note", "One line on what is present or absent. Details are in Findings."),
        ("S", "Sheet: Checks (one row per execution check)"),
        ("#", "Number of the check in the standards workbook."),
        ("Check", "What is verified before a draft leaves Start2 and again before signature."),
        ("Result", "Pass, Fail, N/A or Cannot verify (drop-down)."),
        ("Evidence", "What was seen, with page references."),
        ("S", "Sheet: Documents not supplied (one row per document)"),
        ("Document", "The document as the contract names it, with version or date if given."),
        ("Referenced at", "Contract clause and page where it is named."),
        ("Quote", "The contract text that names or incorporates it."),
        ("Part of contract", "Yes when the contract says the document applies to or forms part of the contract. Those also appear in Findings as type Not supplied."),
        ("Governs", "What the contract says or shows the document covers. Never a guess at its content."),
        ("Effect on this review", "Which findings and coverage rows could not be completed without it."),
        ("Action", "How to obtain it and what to check once it is in hand."),
        ("S", "Rules"),
        ("Counts", "Counts and the to-confirm line use the Findings sheet only. Coverage severity is a view and is not counted."),
        ("Role", "The role is guessed when unclear and recorded on Summary. The only time the skill asks is when a name containing 'Start2' other than Start2 Asia appears: it flags the name and waits for confirmation."),
        ("Not legal advice", "A screen against Start2 standards. Decisions follow the sign-off chain."),
    ]
    header(ws, ["Item", "Description"], [34, 120]); ws.freeze_panes = None
    r = 2
    for a, b in rows:
        if a == "S":
            for j in (1, 2): cell(ws, r, j, b if j == 1 else None, BOLD, SFILL)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        else:
            cell(ws, r, 1, a, BOLD); cell(ws, r, 2, b)
        r += 1


def build(standards, review, out):
    review = review or {}
    S = dict(review.get("summary", {}))
    if review and not S.get("standards_version"):
        S["standards_version"] = standards.get("source", "").replace(".xlsx", "")
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Read Me"; readme(ws)

    # ---- Summary
    ws = wb.create_sheet("Summary"); header(ws, ["Item", "Value", "Where"], [34, 90, 30]); ws.freeze_panes = None
    r = 2; pos = {}
    def section(t):
        nonlocal r
        for j in (1, 2, 3): cell(ws, r, j, t if j == 1 else None, BOLD, SFILL)
        r += 1
    def row(label, key, value=None):
        nonlocal r
        cell(ws, r, 1, label, BOLD); cell(ws, r, 2, value); cell(ws, r, 3)
        pos[key] = r; r += 1
    section("Contract")
    for label, key in (("Contract", "contract"), ("Source file", "file"), ("Status", "status"), ("Contract type", "contract_type"),
                       ("Whose paper", "paper"), ("Standards version", "standards_version")):
        row(label, key, S.get(key))
    section("Our side")
    for label, key in (("Our entity", "our_entity"), ("Our role", "our_role"), ("Role basis", "role_basis"), ("Counterparty", "counterparty")):
        row(label, key, S.get(key))
    dropdown(ws, f"B{pos['status']}", ["Draft", "Executed"]); dropdown(ws, f"B{pos['contract_type']}", ["Service", "Funding", "Vendor", "Other"])
    dropdown(ws, f"B{pos['paper']}", PAPER); dropdown(ws, f"B{pos['our_role']}", ["Provider", "Buyer", "Other"])
    section("Result")
    for k in SEV: row(k, k, f'=COUNTIF(Findings!$B$2:$B${MAXROW},"{k}")')
    ro, c, h = (f"B{pos[x]}" for x in ("our_role", "Critical", "High"))
    row("To confirm", "confirm",
        f'=IF({ro}="","",IF({ro}<>"Provider","Not assessed: the standards cover provider-side contracts only",'
        f'({c}+{h})&" Critical and High findings to confirm"))')
    row("Documents not supplied", "docs", f"=COUNTA('{DOCS}'!$B$2:$B${MAXROW})")
    pp = f"B{pos['paper']}"
    row("Review lead time", "lead", f'=IF({pp}="","",IF({pp}="Start2 template, unmodified","3 business days","5 business days"))')
    row("Sign-off chain", "chain", SIGNOFF)
    for k in SEV: ws.cell(pos[k], 2).alignment = Alignment(horizontal="left", vertical="top")
    colour(ws, f"A{pos['Critical']}:A{pos['Low']}")
    # colour the label cells directly (static) for readability
    for k in SEV: ws.cell(pos[k], 1).fill = PatternFill("solid", fgColor=COLORS[k])
    section("Key facts")
    for f in review.get("facts", []):
        cell(ws, r, 1, f.get("fact"), BOLD); cell(ws, r, 2, f.get("value")); cell(ws, r, 3, f.get("where")); r += 1

    # ---- Findings
    ws = wb.create_sheet("Findings")
    header(ws, ["#", "Severity", "Type", "Clause", "Where", "Quote", "Issue", "Action", "Suggested language", "Subjective"],
           [5, 11, 18, 18, 22, 60, 70, 55, 60, 11])
    order = {k: i for i, k in enumerate(SEV)}
    by = {c["id"]: c for c in standards["clauses"]}
    terms = S.get("defined_terms") or None
    fl = sorted(review.get("findings", []), key=lambda f: order.get(f["severity"], 9))
    for i, f in enumerate(fl, 1):
        vals = [i, f["severity"], f["type"], f["clause"], f.get("where"), rich(f.get("quote")), f.get("issue"), f.get("action"),
                suggestion(f, by, terms), "Yes" if f.get("subjective") else None]
        for j, v in enumerate(vals, 1):
            cell(ws, i + 1, j, v)
    dropdown(ws, f"B2:B{MAXROW}", SEV); dropdown(ws, f"C2:C{MAXROW}", TYPES); dropdown(ws, f"J2:J{MAXROW}", ["Yes"])
    colour(ws, f"B2:B{MAXROW}"); ws.auto_filter.ref = f"A1:J{max(len(fl), 1) + 1}"

    # ---- Coverage
    ws = wb.create_sheet("Coverage")
    header(ws, ["ID", "Clause", "Base severity", "Verdict", "Where", "Severity", "Note"], [14, 36, 12, 14, 26, 12, 70])
    cov = review.get("coverage", {})
    clauses = [x for x in standards["clauses"]]
    for i, x in enumerate(clauses, 2):
        v = cov.get(x["id"], {})
        cell(ws, i, 1, x["id"], BOLD); cell(ws, i, 2, x["clause"]); cell(ws, i, 3, x["base_severity"]); cell(ws, i, 4, v.get("verdict"))
        cell(ws, i, 5, v.get("where")); cell(ws, i, 7, v.get("note"))
        cell(ws, i, 6, f'=IF(D{i}="","",IF(OR(D{i}="Conforming",D{i}="N/A",C{i}="n/a"),"-",IF(D{i}="Partial",'
                       f'INDEX({{"Low","Low","Medium","High"}},MATCH(C{i},{{"Low","Medium","High","Critical"}},0)),C{i})))')
    n = len(clauses) + 1
    dropdown(ws, f"C2:C{n}", SEV + ["n/a"]); dropdown(ws, f"D2:D{n}", VERDICTS)
    colour(ws, f"C2:C{n}"); colour(ws, f"F2:F{n}"); ws.auto_filter.ref = f"A1:G{n}"

    # ---- Checks
    ws = wb.create_sheet("Checks"); header(ws, ["#", "Check", "Result", "Evidence"], [5, 80, 14, 80])
    res = review.get("checks", {})
    for i, ck in enumerate(standards["execution_checks"], 2):
        v = res.get(str(ck["id"]), {})
        label = CHECK_LABEL_OVERRIDE.get(ck["id"], ck["check"])
        cell(ws, i, 1, ck["id"]); cell(ws, i, 2, label); cell(ws, i, 3, v.get("result")); cell(ws, i, 4, v.get("evidence"))
    dropdown(ws, f"C2:C{len(standards['execution_checks']) + 1}", RESULTS)

    # ---- Documents not supplied
    ws = wb.create_sheet(DOCS)
    header(ws, ["#", "Document", "Referenced at", "Quote", "Part of contract", "Governs", "Effect on this review", "Action"],
           [5, 34, 22, 60, 11, 45, 55, 50])
    docs = review.get("documents_not_supplied", [])
    for i, d in enumerate(docs, 1):
        vals = [i, d.get("document"), d.get("referenced_at"), rich(d.get("quote")), "Yes" if d.get("part_of_contract") else "No",
                d.get("governs"), d.get("effect"), d.get("action")]
        for j, v in enumerate(vals, 1):
            cell(ws, i + 1, j, v)
    dropdown(ws, f"E2:E{MAXROW}", ["Yes", "No"]); ws.auto_filter.ref = f"A1:H{max(len(docs), 1) + 1}"
    wb.save(out)


if __name__ == "__main__":
    std = json.load(open(sys.argv[1]))
    if sys.argv[2] == "--template":
        build(std, None, sys.argv[3])
    else:
        build(std, json.load(open(sys.argv[2])), sys.argv[3])
    print("wrote", sys.argv[-1])
