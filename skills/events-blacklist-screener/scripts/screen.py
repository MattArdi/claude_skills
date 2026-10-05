#!/usr/bin/env python3
"""Screen an event guest list against the events blacklist.

Writes a copy of the guest workbook with ONE extra column, "Blacklist Status",
holding OK / Review / Blacklisted. Everything else in the workbook is left as is.

    python screen.py --input guests.xlsx --blacklist Events_Blacklist_Clean.xlsx \
                     --out screened/guests.xlsx [--map map.json] [--dry-run]

Decision rules (see SKILL.md for the reasoning):
    Blacklisted  exact email match, exact LinkedIn profile match, or exact full name
    Review       any other name match (reordered, initials, fuzzy, or a
                 single-token name that is corroborated by company / email domain)
    OK           nothing matched
    (blank)      the row has no name, email or LinkedIn to check
"""
import argparse, csv, json, re, sys, unicodedata
from copy import copy
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

STATUS_HEADER = "Blacklist Status"
OK, REVIEW, BLACKLISTED = "OK", "Review", "Blacklisted"
FUZZY_THRESHOLD = 0.88
FREE_MAIL = {"gmail.com", "googlemail.com", "yahoo.com", "yahoo.com.sg", "yahoo.co.uk", "hotmail.com",
             "hotmail.sg", "outlook.com", "live.com", "msn.com", "icloud.com", "me.com", "gmx.com",
             "gmx.net", "proton.me", "protonmail.com", "qq.com", "163.com", "126.com", "singnet.com.sg"}
COMPANY_NOISE = {"pte", "ltd", "limited", "inc", "llc", "llp", "co", "corp", "corporation", "company",
                 "pl", "gmbh", "sdn", "bhd", "the", "and"}
TITLES = {"mr", "mrs", "ms", "miss", "dr", "prof"}
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+'\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
LI_RE = re.compile(r"linkedin\.com/(?:in|pub)/([^/?#\s]+)", re.I)


# ---------------------------------------------------------------- normalizers
def clean(v):
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip() if v is not None else ""


INVISIBLE = dict.fromkeys(map(ord, "\u200b\u200c\u200d\u200e\u200f\u2060\ufeff\u00ad\u202a\u202b\u202c\u202d\u202e"))
# Cyrillic / Greek letters that look like Latin ones, so a lookalike spelling still compares equal
LOOKALIKE = str.maketrans({
    "а": "a", "в": "b", "е": "e", "к": "k", "м": "m", "н": "h", "о": "o", "р": "p", "с": "c", "т": "t",
    "у": "y", "х": "x", "і": "i", "ј": "j", "ѕ": "s", "һ": "h", "ԁ": "d", "ɡ": "g", "ӏ": "l",
    "α": "a", "β": "b", "ε": "e", "ι": "i", "κ": "k", "ν": "v", "ο": "o", "ρ": "p", "τ": "t", "υ": "u", "χ": "x"})


def norm_text(s):
    s = unicodedata.normalize("NFKD", clean(s).translate(INVISIBLE))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower().translate(LOOKALIKE)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def name_tokens(s):
    t = norm_text(s).split()
    return [w for w in t if w not in TITLES]


LEET = str.maketrans({"0": "o", "1": "i", "3": "e", "4": "a", "5": "s", "7": "t", "@": "a", "$": "s"})
_JOIN_PUNCT = re.compile(r"(?<=\w)[.\u00b7_*'\u2019`](?=\w)")


def alt_forms(raw):
    """Disguised spellings of a name as token lists (beyond the plain normalisation):
    punctuation inside a word ('Go.h'), spaced-out letters ('G o h'), digits for letters ('G0h'),
    stretched letters ('Gooh')."""
    base = name_tokens(raw)
    forms = []
    forms.append(name_tokens(_JOIN_PUNCT.sub("", clean(raw))))
    out, run = [], []
    for t in base + [""]:
        if len(t) == 1:
            run.append(t); continue
        out += ["".join(run)] if len(run) >= 3 else run
        run = []
        if t: out.append(t)
    forms.append(out)
    leet = []
    for t in name_tokens(clean(raw).translate(INVISIBLE)):
        leet.append(t)
    raw_tokens = [w for w in re.split(r"[^\w@$]+", unicodedata.normalize("NFKD", clean(raw)).lower()) if w]
    forms.append([w.translate(LEET) if re.search(r"[a-z]", w) and re.search(r"[0-9@$]", w) else w for w in raw_tokens])
    forms.append([re.sub(r"(.)\1{2,}", r"\1", w) if len(w) > 3 else w for w in base])
    forms.append([re.sub(r"(.)\1", r"\1", w) if len(w) >= 4 else w for w in base])
    seen, res = {tuple(base)}, []
    for f in forms:
        f = [w for w in f if w and w not in TITLES]
        if f and tuple(f) not in seen:
            seen.add(tuple(f)); res.append(f)
    return res


def name_variants(raw):
    """'Eng Hann Lim (registered as E H Lim)' -> ['Eng Hann Lim', 'E H Lim']."""
    raw = clean(raw)
    out = []
    for inner in re.findall(r"\(([^)]*)\)", raw):
        inner = re.sub(r"^(registered as|aka|a\.k\.a\.?|also known as)\s*", "", inner, flags=re.I)
        out.append(inner)
    base = re.sub(r"\([^)]*\)", " ", raw)
    base = re.sub(r"<[^>]*>", " ", base)
    base = EMAIL_RE.sub(" ", base)
    return [v for v in [clean(base)] + [clean(i) for i in out] if v]


def emails_in(v):
    return [e.lower().rstrip(".") for e in EMAIL_RE.findall(clean(v))]


def canon_email(e):
    """Same inbox, different spelling: drop '+tag'; Gmail also ignores dots and treats googlemail.com as gmail.com."""
    local, dom = e.split("@", 1)
    local = local.split("+")[0]
    if dom in ("gmail.com", "googlemail.com"):
        dom, local = "gmail.com", local.replace(".", "")
    return f"{local}@{dom}"


def li_slug(v):
    m = LI_RE.search(clean(v))
    return m.group(1).lower().rstrip("/") if m else None


def company_key(s):
    return " ".join(w for w in norm_text(s).split() if w not in COMPANY_NOISE)


def domain(e):
    return e.split("@", 1)[1]


# ---------------------------------------------------------------- column mapping
def classify_header(h):
    n = norm_text(h)
    if not n:
        return None
    if "linkedin" in n:
        return "linkedin"
    if "mail" in n:
        return "email"
    if re.search(r"\b(also registered|registered as|alias|aka|also known)\b", n):
        return "alias"
    if n in {"first name", "firstname", "first", "given name", "forename", "given names"}:
        return "first"
    if n in {"last name", "lastname", "last", "surname", "family name"}:
        return "last"
    if re.search(r"\b(company|organi[sz]ation|employer|firm|affiliation|institution|org|startup|business)\b", n):
        return "company"
    if "name" in n or n in {"attendee", "guest", "participant", "contact", "person", "invitee", "delegate", "speaker"}:
        return "name"
    return None


def col_idx(ref):
    return column_index_from_string(ref) if isinstance(ref, str) else int(ref)


def infer_mapping(rows, override=None):
    """rows: list of row-value lists (0-based). Returns mapping with 1-based column indexes."""
    if override:
        m = {"header_row": int(override["header_row"])}
        for role in ("name", "first", "last", "company", "linkedin", "alias"):
            if override.get(role):
                m[role] = [col_idx(override[role])] if not isinstance(override[role], list) else [col_idx(x) for x in override[role]]
        m["email"] = [col_idx(x) for x in (override.get("email") or [])]
        m["source"] = "explicit"
        return m
    best = None
    for i, row in enumerate(rows[:25]):
        roles = [(j + 1, classify_header(c)) for j, c in enumerate(row) if isinstance(c, str)]
        roles = [(j, r) for j, r in roles if r]
        kinds = {r for _, r in roles}
        if kinds & {"name", "first", "email", "linkedin"} and (best is None or len(roles) > best[0]):
            best = (len(roles), i, roles)
    if not best:
        return None
    _, i, roles = best
    m = {"header_row": i + 1, "email": [], "source": "headers"}
    for j, r in roles:
        if r == "email":
            m["email"].append(j)
        else:
            m.setdefault(r, []).append(j)
    # content sniffing: headers like "Contact" holding emails / LinkedIn URLs
    data = [r for r in rows[i + 1:i + 201] if any(c not in (None, "") for c in r)]
    used = {j for v in m.values() if isinstance(v, list) for j in v}
    for j in range(1, max((len(r) for r in rows), default=0) + 1):
        if j in used:
            continue
        vals = [clean(r[j - 1]) for r in data if j - 1 < len(r) and clean(r[j - 1])]
        if len(vals) < 1:
            continue
        if sum("@" in v for v in vals) / len(vals) >= 0.5:
            m["email"].append(j); m["source"] += "+sniffed"
        elif sum(bool(li_slug(v)) for v in vals) / len(vals) >= 0.5 and "linkedin" not in m:
            m["linkedin"] = [j]; m["source"] += "+sniffed"
    return m


def describe(m, header):
    def h(j): return f"{get_column_letter(j)}:{clean(header[j-1]) if j-1 < len(header) else ''}"
    return {k: [h(j) for j in v] if isinstance(v, list) else v for k, v in m.items()}


def row_record(row, m):
    def vals(role): return [row[j - 1] for j in m.get(role, []) if j - 1 < len(row)]
    names = []
    first, last = " ".join(clean(v) for v in vals("first")), " ".join(clean(v) for v in vals("last"))
    if first or last:
        names.append(clean(f"{first} {last}"))
    for v in vals("name"):
        names += name_variants(v)
    for v in vals("alias"):
        for part in re.split(r"[;/|]", clean(v)):
            names += name_variants(part)
    emails, slugs = [], []
    for v in vals("email") + vals("name") + vals("company"):
        emails += emails_in(v)
    for v in vals("linkedin") + vals("name"):
        s = li_slug(v)
        if s: slugs.append(s)
    return dict(names=[n for n in dict.fromkeys(names) if name_tokens(n)],
                emails=list(dict.fromkeys(emails)), slugs=list(dict.fromkeys(slugs)),
                company=" ".join(clean(v) for v in vals("company")))


# ---------------------------------------------------------------- matching
def name_match_base(gt, bt):
    """Token lists -> match kind or None. 'single' means one side is a lone token."""
    if not gt or not bt:
        return None
    if len(gt) == 1 or len(bt) == 1:
        lone, other = (gt, bt) if len(gt) == 1 else (bt, gt)
        return "single" if lone[0] in other else None
    if sorted(gt) == sorted(bt):
        return "exact" if gt == bt else "reordered"
    short, long_ = (gt, bt) if len(gt) <= len(bt) else (bt, gt)
    pool, full_hits = list(long_), 0
    initial_used = False
    for t in sorted(short, key=lambda x: -len(x)):
        hit = next((p for p in pool if p == t), None)
        if hit:
            full_hits += 1
        else:
            hit = next((p for p in pool if (len(t) == 1 and p.startswith(t)) or (len(p) == 1 and t.startswith(p))), None)
            initial_used = initial_used or bool(hit)
        if not hit:
            break
        pool.remove(hit)
    else:
        if full_hits >= 1:
            return "initials" if initial_used else "extra-word"
    if SequenceMatcher(None, " ".join(sorted(gt)), " ".join(sorted(bt))).ratio() >= FUZZY_THRESHOLD:
        return "fuzzy"
    return None


def spelling_variants(t):
    """Alternative tokenisations: fused initials split apart ('wh goh' -> 'w h goh'), adjacent tokens joined
    ('wee hong goh' -> 'weehong goh')."""
    out = []
    split = [x for tok in t for x in (list(tok) if 2 <= len(tok) <= 3 and not re.search("[aeiouy]", tok) else [tok])]
    if split != t:
        out.append(split)
    for i in range(len(t) - 1):
        out.append(t[:i] + [t[i] + t[i + 1]] + t[i + 2:])
    return [v for v in out if len(v) >= 2]


def name_match(gt, bt):
    kind = name_match_base(gt, bt)
    if kind:
        return kind
    for g2, b2 in [(g, bt) for g in spelling_variants(gt)] + [(gt, b) for b in spelling_variants(bt)]:
        if name_match_base(g2, b2) in ("exact", "reordered", "initials", "extra-word"):
            return "variant"
    return None


def corroborated(g, b):
    gk, bk = company_key(g["company"]), company_key(b["company"])
    if gk and bk and len(min(gk, bk, key=len)) >= 4 and (gk in bk or bk in gk):
        return "same company"
    gd = {domain(e) for e in g["emails"]} - FREE_MAIL
    bd = {domain(e) for e in b["emails"]} - FREE_MAIL
    if gd & bd:
        return "same email domain"
    return None


def screen_record(g, blacklist, email_idx, slug_idx):
    for e in g["emails"]:
        hit = email_idx.get(canon_email(e))
        if hit:
            note = "" if e in hit["emails"] else " (same inbox after ignoring +tag / Gmail dots)"
            return BLACKLISTED, f"email {e} = blacklist '{hit['display']}'{note}"
    for s in g["slugs"]:
        if s in slug_idx:
            return BLACKLISTED, f"LinkedIn /{s} = blacklist '{slug_idx[s]['display']}'"
    best = None
    rank = {"exact": 0, "reordered": 1, "initials": 2, "extra-word": 3, "variant": 4, "fuzzy": 5, "single": 6}
    for b in blacklist:
        for gn in g["names"]:
            for bn in b["names"]:
                kind = name_match(name_tokens(gn), name_tokens(bn))
                if kind is None:
                    # disguised spellings: always Review, never an automatic Blacklisted
                    alt = [name_match(g2, name_tokens(bn)) for g2 in alt_forms(gn)] + \
                          [name_match(name_tokens(gn), b2) for b2 in alt_forms(bn)]
                    kind = "variant" if any(k and k != "single" for k in alt) else None
                if kind is None:
                    continue
                if kind == "exact":
                    return BLACKLISTED, f"name '{gn}' = blacklist '{b['display']}' (exact name)"
                if kind == "single":
                    why = corroborated(g, b)
                    if not why:
                        continue
                    kind = f"single-name + {why}"
                key = rank.get(kind.split(" ")[0], 5)
                if best is None or key < best[0]:
                    best = (key, f"name '{gn}' ~ blacklist '{b['display']}' ({kind})", b)
    if best:
        return REVIEW, best[1]
    return OK, ""


# ---------------------------------------------------------------- IO
def sheet_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def load_blacklist(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = [wb["Blacklist"]] if "Blacklist" in wb.sheetnames else wb.worksheets
    for ws in sheets:
        rows = sheet_rows(ws)
        m = infer_mapping(rows)
        if m and (m.get("name") or m.get("first")):
            hdr = m["header_row"]
            entries = []
            for row in rows[hdr:]:
                r = row_record(row, m)
                if not (r["names"] or r["emails"] or r["slugs"]):
                    continue
                r["display"] = r["names"][0] if r["names"] else r["emails"][0]
                entries.append(r)
            email_idx = {canon_email(e): b for b in entries for e in b["emails"]}
            slug_idx = {s: b for b in entries for s in b["slugs"]}
            return entries, email_idx, slug_idx, ws.title
    sys.exit("Could not find a Name column in the blacklist workbook.")


def open_guests(path):
    p = Path(path)
    if p.suffix.lower() == ".csv":
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = p.stem[:31]
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = p.read_text(encoding=enc); break
            except UnicodeDecodeError:
                continue
        for row in csv.reader(text.splitlines()):
            ws.append(row)
        return wb, wb
    if p.suffix.lower() == ".xls":
        sys.exit("Legacy .xls is not supported. Convert it to .xlsx first.")
    keep = p.suffix.lower() == ".xlsm"
    return openpyxl.load_workbook(p, keep_vba=keep), openpyxl.load_workbook(p, data_only=True, keep_vba=keep)


def copy_style(src, dst):
    if src.has_style:
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.protection = copy(src.alignment), copy(src.protection)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", required=True)
    ap.add_argument("--blacklist", required=True)
    ap.add_argument("--out", help="output .xlsx path (omit with --dry-run)")
    ap.add_argument("--map", help="JSON file or string overriding column mapping; "
                    '{"header_row":3,"name":"B","email":["E","F"]} or keyed by sheet name')
    ap.add_argument("--dry-run", action="store_true", help="print mapping and matches, write nothing")
    a = ap.parse_args()
    if not a.dry_run and not a.out:
        ap.error("--out is required unless --dry-run")

    bl, email_idx, slug_idx, bl_sheet = load_blacklist(a.blacklist)
    wb, wbv = open_guests(a.input)
    override = None
    if a.map:
        override = json.loads(Path(a.map).read_text() if Path(a.map).exists() else a.map)

    report = dict(blacklist_entries=len(bl), sheets={}, flagged=[], warnings=[])
    for ws in wb.worksheets:
        wsv = wbv[ws.title]
        rows = sheet_rows(wsv)
        ov = None
        if override:
            ov = override.get(ws.title) if ws.title in override else (override if "header_row" in override else None)
        m = infer_mapping(rows, ov)
        if not m or not (m.get("name") or m.get("first") or m["email"] or m.get("linkedin")):
            report["sheets"][ws.title] = "skipped: no name/email/LinkedIn column found (pass --map to force)"
            continue
        if getattr(ws, "_charts", None) or getattr(ws, "_images", None) or ws.tables:
            report["warnings"].append(f"{ws.title}: charts/images/tables may not survive the save; check the output")
        hdr = m["header_row"]
        status_col = next((c.column for c in ws[hdr] if clean(c.value).lower() == STATUS_HEADER.lower()), None)
        if status_col is None:
            # first free column after the last one holding any value (sheets often carry empty formatted columns)
            status_col = max((j + 1 for row in rows for j, c in enumerate(row) if c not in (None, "")), default=0) + 1
        hc = ws.cell(hdr, status_col, STATUS_HEADER)
        left = ws.cell(hdr, max(status_col - 1, 1))
        copy_style(left, hc)
        ws.column_dimensions[get_column_letter(status_col)].width = 18
        if ws.auto_filter.ref:
            c1, r1, c2, r2 = range_boundaries(ws.auto_filter.ref)
            if c2 < status_col:
                ws.auto_filter.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(status_col)}{r2}"
        counts = {OK: 0, REVIEW: 0, BLACKLISTED: 0, "unchecked": 0}
        for r in range(hdr + 1, ws.max_row + 1):
            row = rows[r - 1] if r - 1 < len(rows) else []
            if not any(c not in (None, "") for c in row):
                continue
            rec = row_record(row, m)
            cell = ws.cell(r, status_col)
            copy_style(ws.cell(r, max(status_col - 1, 1)), cell)
            if not (rec["names"] or rec["emails"] or rec["slugs"]):
                counts["unchecked"] += 1
                continue
            status, why = screen_record(rec, bl, email_idx, slug_idx)
            cell.value = status
            counts[status] += 1
            if status != OK:
                report["flagged"].append(dict(sheet=ws.title, row=r, guest=rec["names"][0] if rec["names"] else rec["emails"][0],
                                              status=status, why=why))
        header = rows[hdr - 1]
        report["sheets"][ws.title] = dict(mapping=describe(m, header), counts=counts, status_column=get_column_letter(status_col))
    if not a.dry_run:
        Path(a.out).parent.mkdir(parents=True, exist_ok=True)
        wb.save(a.out)
        report["output"] = a.out
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
