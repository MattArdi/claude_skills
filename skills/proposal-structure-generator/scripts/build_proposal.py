"""Build a proposal skeleton from a filled-in checklist.

Usage:
  python build_proposal.py --checklist checklist.xlsx --blocks blocks.json -o out.docx
        [--template ../assets/Start2_Proposal_Template_for_Public_Partners.docx]
        [--map ../assets/section_map_public_partners.json]

blocks.json maps checklist IDs to Notion block content fetched at build time:
  {"PP-008": {"status": "Reviewed", "text": "<Notion rich-text markdown>"}, ...}
Status in blocks.json overrides the snapshot status in the section map.

The script starts from the master template, keeps or removes each heading per the
checklist, inserts Notion text for Reviewed blocks, writes "Subheader was empty."
where a heading has no content, and applies the writing guide rules that can be
enforced in code (palette values, Inter font, heading typography, list numbering,
wording fixes).
"""
import argparse, copy, json, os, re, sys
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_result import results, find_header  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ASSETS = os.path.join(HERE, "..", "assets")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
EMPTY_NOTE = "Subheader was empty."
LEVELS = {"H1": 1, "H2": 2, "H3": 3, "H4": 4}
# Template color values (most frequent value across both master templates)
PALETTE = {"52389D": "51389D", "53389E": "51389D", "4C2E71": "51389D",
           "F1EFFD": "F2EFFE", "F2EFFF": "F2EFFE", "EEF3FC": "EFF4FC", "919191": "888888"}
WORDING = [(r"\bProgrammes\b", "Programs"), (r"\bProgramme\b", "Program"), (r"\bprogrammes\b", "programs"),
           (r"\bprogramme\b", "program"), (r"\bBMWK\b", "BMWE"), (r"\s*[—–]\s*", ", "),
           (r"\bStart2(?! Group)(?! Asia)\b", "Start2 Group")]
BULLET_NUM_ID = "5"      # template bullet list (abstract 12)
DECIMAL_ABSTRACT = "8"   # template "1." list used by Key Benefits


def w(tag):
    return qn("w:" + tag)


def ptext(p):
    return "".join(t.text or "" for t in p.iter(w("t"))).strip()


def pstyle(el):
    if el.tag != w("p"):
        return None
    s = el.find("./w:pPr/w:pStyle", {"w": W})
    return s.get(w("val")) if s is not None else None


def heading_level(el):
    s = pstyle(el)
    m = re.match(r"Heading(\d)$", s or "")
    return int(m.group(1)) if m else None


def norm(s):
    return re.sub(r"\s+", " ", s or "").strip().lower()


def is_placeholder(el):
    return el.tag == w("p") and re.match(r"^\[.*\]$", ptext(el)) is not None


# ---------- Notion markdown to runs ----------

def tokenize(md):
    """Notion rich-text markdown -> list of (text, bold, italic); '\n' marks a line break."""
    md = re.sub(r"(?<!\\)\[([^\]]+)\]\((?:https?://)[^)]*\)", r"\1", md)  # drop links, keep text
    md = md.replace("<br>", "\n")
    out, buf, b, i, k = [], [], False, False, 0

    def flush():
        if buf:
            out.append(("".join(buf), b, i))
            buf.clear()
    while k < len(md):
        ch = md[k]
        if ch == "\\" and k + 1 < len(md):
            buf.append(md[k + 1]); k += 2; continue
        if md.startswith("**", k):
            flush(); b = not b; k += 2; continue
        if ch in "*_":
            flush(); i = not i; k += 1; continue
        if ch == "\n":
            flush(); out.append(("\n", b, i)); k += 1; continue
        buf.append(ch); k += 1
    flush()
    return out


def to_blocks(md):
    """-> list of paragraphs: dict(kind='p'|'bullet'|'num', lines=[[runs], ...])."""
    lines, cur = [], []
    for t in tokenize(md):
        if t[0] == "\n":
            lines.append(cur); cur = []
        else:
            cur.append(t)
    lines.append(cur)
    paras, group = [], None
    for runs in lines:
        plain = "".join(r[0] for r in runs)
        if not plain.strip():
            group = None
            continue
        m = re.match(r"^\s*(\d+\.|[•\-])\s*", plain)
        if m:
            kind = "num" if m.group(1)[0].isdigit() else "bullet"
            start = int(m.group(1)[:-1]) if kind == "num" else None
            runs = strip_prefix(runs, len(m.group(0)))
            if not "".join(r[0] for r in runs).strip():
                continue
            paras.append(dict(kind=kind, lines=[runs], start=start))
            group = None
            continue
        runs = strip_prefix(runs, len(plain) - len(plain.lstrip()))
        if group is None:
            group = dict(kind="p", lines=[])
            paras.append(group)
        group["lines"].append(runs)
    return paras


def strip_prefix(runs, n):
    out = []
    for text, b, i in runs:
        if n >= len(text):
            n -= len(text); continue
        out.append((text[n:], b, i)); n = 0
    if out:
        out[-1] = (out[-1][0].rstrip(), out[-1][1], out[-1][2])
    return out


# ---------- XML builders ----------

class Builder:
    def __init__(self, doc):
        self.doc = doc
        self.numbering = doc.part.numbering_part.element
        ids = [int(n.get(w("numId"))) for n in self.numbering.findall(w("num"))]
        self.next_num = max(ids) + 1

    def new_decimal_list(self):
        num = etree.Element(w("num"))
        self.numbering.findall(w("num"))[-1].addnext(num)  # schema: num entries precede numIdMacAtCleanup
        num.set(w("numId"), str(self.next_num))
        a = etree.SubElement(num, w("abstractNumId")); a.set(w("val"), DECIMAL_ABSTRACT)
        o = etree.SubElement(num, w("lvlOverride")); o.set(w("ilvl"), "0")
        s = etree.SubElement(o, w("startOverride")); s.set(w("val"), "1")
        self.next_num += 1
        return str(self.next_num - 1)

    @staticmethod
    def run(text, bold=False, italic=False, color=None):
        r = etree.Element(w("r"))
        if bold or italic or color:
            rpr = etree.SubElement(r, w("rPr"))
            if bold:
                etree.SubElement(rpr, w("b")); etree.SubElement(rpr, w("bCs"))
            if italic:
                etree.SubElement(rpr, w("i")); etree.SubElement(rpr, w("iCs"))
            if color:
                c = etree.SubElement(rpr, w("color")); c.set(w("val"), color)
        t = etree.SubElement(r, w("t")); t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return r

    def para(self, lines, style=None, num_id=None, after="240"):
        p = etree.Element(w("p"))
        ppr = etree.SubElement(p, w("pPr"))
        if style:
            s = etree.SubElement(ppr, w("pStyle")); s.set(w("val"), style)
        if num_id:
            npr = etree.SubElement(ppr, w("numPr"))
            il = etree.SubElement(npr, w("ilvl")); il.set(w("val"), "0")
            ni = etree.SubElement(npr, w("numId")); ni.set(w("val"), num_id)
        sp = etree.SubElement(ppr, w("spacing")); sp.set(w("after"), after)
        for li, runs in enumerate(lines):
            if li:
                br = etree.Element(w("r")); etree.SubElement(br, w("br")); p.append(br)
            for text, b, i in runs:
                p.append(self.run(text, b, i))
        return p

    def notion_paragraphs(self, md):
        out, list_id = [], None
        for blk in to_blocks(md):
            if blk["kind"] == "num":
                # a typed "1." starts a new list; a higher number continues the current one
                if list_id is None or blk.get("start") == 1:
                    list_id = self.new_decimal_list()
                out.append(self.para(blk["lines"], "ListParagraph", list_id, after="120"))
            elif blk["kind"] == "bullet":
                out.append(self.para(blk["lines"], "ListParagraph", BULLET_NUM_ID, after="120"))
            else:
                out.append(self.para(blk["lines"]))
        return out

    def note(self):
        return self.para([[(EMPTY_NOTE, False, True)]])

    def heading(self, text, level):
        p = etree.Element(w("p"))
        ppr = etree.SubElement(p, w("pPr"))
        s = etree.SubElement(ppr, w("pStyle")); s.set(w("val"), f"Heading{level}")
        p.append(self.run(text))
        return p


# ---------- template sections ----------

def template_sections(body):
    """Split body into sections: {norm(heading text): dict(el, level, content=[...])}."""
    secs, order, cur = {}, [], None
    for el in list(body):
        if el.tag == w("sectPr"):
            continue
        lv = heading_level(el)
        if lv:
            cur = dict(el=el, level=lv, text=ptext(el), content=[])
            if cur["text"]:
                secs[norm(cur["text"])] = cur
            order.append(cur)
        elif cur is not None:
            cur["content"].append(el)
    return secs, order


def section_end(head):
    """Last element belonging to the section that starts at head."""
    el = head
    while True:
        nxt = el.getnext()
        if nxt is None or nxt.tag == w("sectPr") or heading_level(nxt):
            return el
        el = nxt


def insert_after(anchor, els):
    for el in els:
        anchor.addnext(el)
        anchor = el
    return anchor


def build(checklist, blocks, smap, template, out):
    wb = load_workbook(checklist)
    ws = wb.worksheets[0]
    hr, cols = find_header(ws)
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        rid = ws.cell(r, cols["ID"]).value
        if rid:
            rows.append(dict(id=str(rid).strip(), level=str(ws.cell(r, cols["Level"]).value or "").strip(),
                             section=str(ws.cell(r, cols["Section"]).value or "").strip(),
                             include=ws.cell(r, cols["Include"]).value or "No"))
    # live Notion status overrides the snapshot
    live = {k: dict(v) for k, v in smap.items()}
    for k, v in blocks.items():
        if k in live and v.get("status") and live[k]["notion_status"] != "Not used":
            live[k]["notion_status"] = v["status"]
    res = results(rows, live)

    doc = Document(template)
    body = doc.element.body
    B = Builder(doc)
    secs, order = template_sections(body)
    log = []

    # blank headings in the template produce empty TOC entries
    for s in order:
        if not s["text"]:
            s["el"].getparent().remove(s["el"])
    toc = secs.get(norm("Table of Contents"))
    cursor = toc["content"][-1] if toc and toc["content"] else (toc["el"] if toc else body[0])

    i = 0
    while i < len(rows):
        r, result = rows[i], res[i]
        info = live.get(r["id"], {})
        if r["level"] == "Fixed":
            i += 1; continue
        lv = LEVELS.get(r["level"])
        if r["level"] == "Body":
            i += 1; continue  # handled with its parent heading
        sec = secs.get(norm(r["section"]))
        if result.startswith("Omitted"):
            if sec:
                for el in [sec["el"]] + sec["content"]:
                    if el.getparent() is not None:
                        el.getparent().remove(el)
            log.append((r["id"], r["section"], "removed"))
            i += 1; continue
        if sec is None:  # heading the template does not have
            head = B.heading(r["section"], lv)
            cursor = insert_after(cursor, [head])
            content = []
            log.append((r["id"], r["section"], "heading added"))
        else:
            head, content = sec["el"], sec["content"]
        # decide content
        new = []
        if info.get("template_content") == "Replaced":
            for el in content:
                el.getparent().remove(el)
            content = []
            j = i + 1
            while j < len(rows) and rows[j]["level"] == "Body":
                if res[j].startswith("Notion text"):
                    new += B.notion_paragraphs(blocks.get(rows[j]["id"], {}).get("text", ""))
                elif not res[j].startswith("Omitted"):
                    new.append(B.note())
                j += 1
            insert_after(head, new)
            cursor = section_end(head)
            i = j
            log.append((r["id"], r["section"], f"template text replaced by {len(new)} Notion paragraphs"))
            continue
        if result.startswith("Heading + Notion text"):
            md = blocks.get(r["id"], {}).get("text", "")
            new = B.notion_paragraphs(md)
            drop = [el for el in content if el.tag == w("p")]
        elif "note:" in result:
            new = [B.note()]
            drop = [el for el in content if is_placeholder(el)]
        else:
            drop = [el for el in content if is_placeholder(el)]
        # insert at the first dropped paragraph, else directly under the heading
        anchor = head
        if drop and "note:" not in result:
            prev = drop[0].getprevious()
            anchor = prev if prev is not None else head
        for el in drop:
            el.getparent().remove(el)
        insert_after(anchor, new)
        cursor = section_end(head)
        log.append((r["id"], r["section"], result))
        i += 1

    apply_guide(doc)
    doc.save(out)
    return log


def apply_guide(doc):
    """Writing-guide rules enforced in code."""
    body = doc.element.body
    parts = [doc.part.element, doc.styles.element] + [s.footer._element for s in doc.sections] + \
            [s.header._element for s in doc.sections]
    for root in parts:
        for el in root.iter():
            for att in (w("fill"), w("color"), w("val")):
                v = el.get(att)
                if v and v.upper() in PALETTE:
                    el.set(att, PALETTE[v.upper()])
    # wording rules on body text
    for t in body.iter(w("t")):
        if t.text:
            s = t.text
            for pat, rep in WORDING:
                s = re.sub(pat, rep, s)
            t.text = s
    st = doc.styles.element
    # Inter everywhere, Arial fallback handled by the recipient's system
    rfd = st.find(".//w:docDefaults/w:rPrDefault/w:rPr/w:rFonts", {"w": W})
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfd.set(w(a), "Inter")
    for s in st.findall(w("style")):
        rpr = s.find(w("rPr"))
        if rpr is not None:
            f = rpr.find(w("rFonts"))
            if f is not None:
                for a in ("ascii", "hAnsi", "eastAsia", "cs"):
                    if f.get(w(a)):
                        f.set(w(a), "Inter")
                for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                    f.attrib.pop(w(a), None)
    for r in body.iter(w("rFonts")):
        for a in ("ascii", "hAnsi"):
            if r.get(w(a)):
                r.set(w(a), "Inter")
    # heading typography from the guide
    spec = {"Heading1": ("51389D", 40), "Heading2": ("374050", 32), "Heading3": ("374050", 28), "Heading4": ("374050", 24)}
    for s in st.findall(w("style")):
        sid = s.get(w("styleId"))
        if sid in spec:
            color, sz = spec[sid]
            rpr = s.find(w("rPr"))
            if rpr is None:
                rpr = etree.SubElement(s, w("rPr"))
            for tag in ("b", "color", "sz", "szCs"):
                for x in rpr.findall(w(tag)):
                    rpr.remove(x)
            etree.SubElement(rpr, w("b"))
            c = etree.SubElement(rpr, w("color")); c.set(w("val"), color)
            for tag in ("sz", "szCs"):
                x = etree.SubElement(rpr, w(tag)); x.set(w("val"), str(sz))
    # refresh the TOC when the file is opened
    settings = doc.settings.element
    if settings.find(w("updateFields")) is None:
        u = etree.Element(w("updateFields")); u.set(w("val"), "true")
        later = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                 "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                 "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                 "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")
        nxt = next((c for c in settings if etree.QName(c).localname in later), None)
        if nxt is not None:
            nxt.addprevious(u)
        else:
            settings.append(u)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--checklist", required=True)
    ap.add_argument("--blocks", required=True)
    ap.add_argument("--template", default=os.path.join(ASSETS, "Start2_Proposal_Template_for_Public_Partners.docx"))
    ap.add_argument("--map", default=os.path.join(ASSETS, "section_map_public_partners.json"))
    ap.add_argument("-o", "--output", required=True)
    a = ap.parse_args()
    smap = json.load(open(a.map, encoding="utf-8"))["sections"]
    blocks = json.load(open(a.blocks, encoding="utf-8"))
    for rid, sec, what in build(a.checklist, blocks, smap, a.template, a.output):
        print(f"{rid}  {sec}: {what}")
    print("->", a.output)


if __name__ == "__main__":
    main()
