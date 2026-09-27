"""Add a Build Result column to a filled-in proposal checklist.

Usage: python build_result.py <checklist.xlsx> [-o output.xlsx] [--map section_map.json]

Input sheet columns: ID | No. | Level | Section | Include (Yes / No / Always).
Build Result says what the proposal build will put under each row, using the
template content and Notion review status recorded in the section map.
"""
import argparse, json, os
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_MAP = os.path.join(HERE, "..", "assets", "section_map_public_partners.json")
EMPTY_NOTE = "Subheader was empty."
LEVELS = {"H1": 1, "H2": 2, "H3": 3, "H4": 4}


def find_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=20):
        vals = [str(c.value).strip() if c.value is not None else "" for c in row]
        if "ID" in vals and "Include" in vals:
            return row[0].row, {v: i + 1 for i, v in enumerate(vals) if v}
    raise SystemExit("No header row with ID and Include found.")


def results(rows, smap):
    """rows: list of dicts with id, level, include. Returns Build Result per row, in order."""
    # included = the row is on, and every ancestor heading is on
    stack, included = [], []
    for r in rows:
        on = str(r["include"]).strip().lower() in ("yes", "always")
        lv = LEVELS.get(r["level"])
        if lv:
            while stack and stack[-1][0] >= lv:
                stack.pop()
            parent_on = all(s[1] for s in stack)
            stack.append((lv, on and parent_on))
            included.append(on and parent_on)
        else:  # Fixed or Body rows follow their own flag and the enclosing heading
            included.append(on and all(s[1] for s in stack))
    out = []
    for i, r in enumerate(rows):
        info = smap.get(r["id"], {})
        tc, st = info.get("template_content", "None"), info.get("notion_status", "No block")
        on = str(r["include"]).strip().lower() in ("yes", "always")
        if not included[i]:
            out.append("Omitted (parent heading excluded)" if on else "Omitted")
            continue
        table = " + template table" if tc == "Table" else ""
        if r["level"] == "Fixed":
            out.append("Kept as in template")
        elif r["level"] == "Body":
            out.append("Notion text (no heading)" if st == "Reviewed" else f"Note: {EMPTY_NOTE}")
        elif st == "Reviewed":
            out.append("Heading + Notion text" + table)
        elif st == "Not used":
            lv = LEVELS[r["level"]]
            kids = False
            for j in range(i + 1, len(rows)):
                lj = LEVELS.get(rows[j]["level"])
                if lj is not None and lj <= lv:
                    break
                kids = kids or included[j]
            if tc == "Table":
                out.append("Heading + template table")
            elif tc == "Fixed text":
                out.append("Heading + template text")
            elif kids:
                out.append("Heading")
            else:
                out.append(f"Heading + note: {EMPTY_NOTE}")
        else:  # Pending Review, Empty, No block
            out.append(f"Heading + note: {EMPTY_NOTE}" + table)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("checklist")
    ap.add_argument("-o", "--output")
    ap.add_argument("--map", default=DEFAULT_MAP)
    a = ap.parse_args()
    smap = json.load(open(a.map, encoding="utf-8"))["sections"]
    wb = load_workbook(a.checklist)
    ws = wb.worksheets[0]
    hr, cols = find_header(ws)
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        rid = ws.cell(r, cols["ID"]).value
        if not rid:
            continue
        rows.append(dict(row=r, id=str(rid).strip(), level=str(ws.cell(r, cols["Level"]).value or "").strip(),
                         include=ws.cell(r, cols["Include"]).value or "No"))
    unknown = [x["id"] for x in rows if x["id"] not in smap]
    if unknown:
        print("Warning: IDs not in section map:", ", ".join(unknown))
    res = results(rows, smap)
    bc = cols.get("Build Result") or (max(cols.values()) + 1)
    src = ws.cell(hr, cols["Include"])
    h = ws.cell(hr, bc, "Build Result")
    h.font, h.fill, h.alignment, h.border = copy(src.font), copy(src.fill), Alignment(vertical="center"), copy(src.border)
    lilac, thin = PatternFill("solid", fgColor="F2EFFE"), Border(bottom=Side(style="thin", color="F2EFFE"))
    for x, v in zip(rows, res):
        c = ws.cell(x["row"], bc, v)
        c.font = Font(name="Arial", size=10, color="888888" if v.startswith("Omitted") else "374050", bold=x["level"] == "H1")
        c.alignment, c.border = Alignment(vertical="center", wrap_text=True), thin
        if x["level"] == "H1":
            c.fill = lilac
    ws.column_dimensions[ws.cell(hr, bc).column_letter].width = 42
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A{hr}:{ws.cell(hr, bc).column_letter}{ws.max_row}"
    out = a.output or a.checklist.replace(".xlsx", "_build_result.xlsx")
    wb.save(out)
    print(f"{len(rows)} rows -> {out}")


if __name__ == "__main__":
    main()
